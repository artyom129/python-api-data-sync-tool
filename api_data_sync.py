from __future__ import annotations

import csv
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


API_URL = "https://jsonplaceholder.typicode.com/users"
SAMPLE_FILE = Path("sample_api_response.json")
OUTPUT_DIR = Path("output")
DATABASE_FILE = OUTPUT_DIR / "records.db"
CSV_FILE = OUTPUT_DIR / "api_records_clean.csv"
EXCEL_FILE = OUTPUT_DIR / "api_records_report.xlsx"

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def load_sample_data(path: Path = SAMPLE_FILE) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Sample file not found: {path}")

    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if not isinstance(data, list):
        raise ValueError("Sample JSON must contain a list of objects")

    return data


def fetch_api_data(
    use_sample: bool = False,
    api_url: str = API_URL,
    timeout_seconds: int = 15,
) -> list[dict[str, Any]]:
    """Fetch records from the API and fall back to local demo data on failure."""
    if use_sample:
        return load_sample_data()

    try:
        response = requests.get(api_url, timeout=timeout_seconds)
        response.raise_for_status()
        data = response.json()
        if not isinstance(data, list):
            raise ValueError("API response must contain a list of objects")
        return data
    except (requests.RequestException, ValueError) as error:
        print(f"API request failed; using sample data instead: {error}")
        return load_sample_data()


def normalize_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_email(value: Any) -> str:
    return normalize_text(value).lower()


def extract_nested_text(record: dict[str, Any], field: str, nested_key: str) -> str:
    value = record.get(field, "")
    if isinstance(value, dict):
        return normalize_text(value.get(nested_key, ""))
    return normalize_text(value)


def clean_records(
    raw_records: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    clean: list[dict[str, Any]] = []
    invalid: list[dict[str, Any]] = []
    seen_emails: set[str] = set()
    synced_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    for record in raw_records:
        cleaned = {
            "source_id": normalize_text(record.get("id")),
            "name": normalize_text(record.get("name")).title(),
            "username": normalize_text(record.get("username")),
            "email": normalize_email(record.get("email")),
            "phone": normalize_text(record.get("phone")),
            "website": normalize_text(record.get("website")),
            "city": extract_nested_text(record, "address", "city"),
            "company": extract_nested_text(record, "company", "name"),
            "synced_at": synced_at,
            "error": "",
        }

        errors: list[str] = []
        if not cleaned["source_id"]:
            errors.append("Missing source id")
        if not cleaned["name"]:
            errors.append("Missing name")
        if not EMAIL_RE.match(cleaned["email"]):
            errors.append("Invalid email")
        elif cleaned["email"] in seen_emails:
            errors.append("Duplicate email")

        if errors:
            cleaned["error"] = "; ".join(errors)
            invalid.append(cleaned)
            continue

        seen_emails.add(cleaned["email"])
        clean.append(cleaned)

    return clean, invalid


def save_to_sqlite(
    records: list[dict[str, Any]],
    database_file: Path = DATABASE_FILE,
) -> None:
    database_file.parent.mkdir(parents=True, exist_ok=True)

    with sqlite3.connect(database_file) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS api_records (
                source_id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                username TEXT,
                email TEXT UNIQUE NOT NULL,
                phone TEXT,
                website TEXT,
                city TEXT,
                company TEXT,
                synced_at TEXT NOT NULL
            )
            """
        )

        connection.executemany(
            """
            INSERT INTO api_records (
                source_id, name, username, email, phone,
                website, city, company, synced_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(source_id) DO UPDATE SET
                name = excluded.name,
                username = excluded.username,
                email = excluded.email,
                phone = excluded.phone,
                website = excluded.website,
                city = excluded.city,
                company = excluded.company,
                synced_at = excluded.synced_at
            """,
            [
                (
                    record["source_id"],
                    record["name"],
                    record["username"],
                    record["email"],
                    record["phone"],
                    record["website"],
                    record["city"],
                    record["company"],
                    record["synced_at"],
                )
                for record in records
            ],
        )
        connection.commit()


def export_to_csv(
    records: list[dict[str, Any]],
    output_file: Path = CSV_FILE,
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "source_id",
        "name",
        "username",
        "email",
        "phone",
        "website",
        "city",
        "company",
        "synced_at",
    ]

    with output_file.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(
            {field: record.get(field, "") for field in fieldnames}
            for record in records
        )


def style_header(worksheet, row: int = 1) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def write_records_sheet(worksheet, records: list[dict[str, Any]], include_error: bool) -> None:
    headers = [
        "source_id",
        "name",
        "username",
        "email",
        "phone",
        "website",
        "city",
        "company",
        "synced_at",
    ]
    if include_error:
        headers.append("error")

    worksheet.append(headers)
    for record in records:
        worksheet.append([record.get(header, "") for header in headers])

    style_header(worksheet)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column in worksheet.columns:
        column_letter = column[0].column_letter
        max_length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[column_letter].width = min(max_length + 3, 40)


def create_dashboard_sheet(
    worksheet,
    clean_records: list[dict[str, Any]],
    invalid_records: list[dict[str, Any]],
) -> None:
    metrics = [
        ("Raw records", len(clean_records) + len(invalid_records)),
        ("Clean records", len(clean_records)),
        ("Invalid records", len(invalid_records)),
        ("Unique companies", len({r["company"] for r in clean_records if r["company"]})),
        ("Unique cities", len({r["city"] for r in clean_records if r["city"]})),
    ]

    worksheet["A1"] = "API Data Sync Report"
    worksheet["A1"].font = Font(size=18, bold=True, color="1F4E78")
    worksheet.merge_cells("A1:B1")
    worksheet.append(["Metric", "Value"])
    for metric, value in metrics:
        worksheet.append([metric, value])

    style_header(worksheet, row=2)
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 18

    chart = BarChart()
    chart.title = "Sync Summary"
    chart.y_axis.title = "Count"
    data = Reference(worksheet, min_col=2, min_row=3, max_row=7)
    categories = Reference(worksheet, min_col=1, min_row=3, max_row=7)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 14
    worksheet.add_chart(chart, "D2")


def export_to_excel(
    clean_records: list[dict[str, Any]],
    invalid_records: list[dict[str, Any]],
    output_file: Path = EXCEL_FILE,
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    clean_sheet = workbook.create_sheet("Clean Records")
    invalid_sheet = workbook.create_sheet("Invalid Records")

    create_dashboard_sheet(dashboard, clean_records, invalid_records)
    write_records_sheet(clean_sheet, clean_records, include_error=False)
    write_records_sheet(invalid_sheet, invalid_records, include_error=True)
    workbook.save(output_file)


def main() -> None:
    raw_records = fetch_api_data()
    clean, invalid = clean_records(raw_records)

    save_to_sqlite(clean)
    export_to_csv(clean)
    export_to_excel(clean, invalid)

    print("API data sync complete")
    print(f"Clean records: {len(clean)}")
    print(f"Invalid records: {len(invalid)}")
    print(f"SQLite database: {DATABASE_FILE}")
    print(f"CSV export: {CSV_FILE}")
    print(f"Excel report: {EXCEL_FILE}")


if __name__ == "__main__":
    main()
