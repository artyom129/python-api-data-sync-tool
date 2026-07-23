from __future__ import annotations

import csv
import json
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference


API_URL = "https://jsonplaceholder.typicode.com/users"
SAMPLE_FILE = Path("sample_api_response.json")

OUTPUT_DIR = Path("output")
DATABASE_FILE = OUTPUT_DIR / "records.db"
CSV_FILE = OUTPUT_DIR / "api_records_clean.csv"
EXCEL_FILE = OUTPUT_DIR / "api_records_report.xlsx"

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def fetch_api_data(use_sample: bool = False) -> list[dict[str, Any]]:
    """
    Fetch records from an API endpoint.

    If the API is unavailable or use_sample=True, the script uses
    sample_api_response.json so the project can run offline.
    """
    if use_sample:
        return load_sample_data()

    try:
        response = requests.get(API_URL, timeout=15)
        response.raise_for_status()
        data = response.json()

        if not isinstance(data, list):
            raise ValueError("API response must be a list of objects")

        return data

    except Exception as error:
        print(f"API request failed, using sample data instead: {error}")
        return load_sample_data()


def load_sample_data() -> list[dict[str, Any]]:
    if not SAMPLE_FILE.exists():
        raise FileNotFoundError(f"Sample file not found: {SAMPLE_FILE}")

    with SAMPLE_FILE.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if not isinstance(data, list):
        raise ValueError("sample_api_response.json must contain a list")

    return data


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_email(value: Any) -> str:
    return normalize_text(value).lower()


def extract_company_name(record: dict[str, Any]) -> str:
    company = record.get("company", "")
    if isinstance(company, dict):
        return normalize_text(company.get("name", ""))
    return normalize_text(company)


def extract_city(record: dict[str, Any]) -> str:
    address = record.get("address", "")
    if isinstance(address, dict):
        return normalize_text(address.get("city", ""))
    return normalize_text(address)


def clean_records(raw_records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    clean: list[dict[str, Any]] = []
    invalid: list[dict[str, Any]] = []
    seen_emails: set[str] = set()

    for record in raw_records:
        cleaned = {
            "source_id": normalize_text(record.get("id")),
            "name": normalize_text(record.get("name")).title(),
            "username": normalize_text(record.get("username")),
            "email": normalize_email(record.get("email")),
            "phone": normalize_text(record.get("phone")),
            "website": normalize_text(record.get("website")),
            "city": extract_city(record),
            "company": extract_company_name(record),
            "synced_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "error": "",
        }

        errors: list[str] = []

        if not cleaned["source_id"]:
            errors.append("Missing source id")
        if not cleaned["name"]:
            errors.append("Missing name")
        if not EMAIL_RE.match(cleaned["email"]):
            errors.append("Invalid email")
        if cleaned["email"] in seen_emails:
            errors.append("Duplicate email")

        if errors:
            cleaned["error"] = "; ".join(errors)
            invalid.append(cleaned)
            continue

        seen_emails.add(cleaned["email"])
        clean.append(cleaned)

    return clean, invalid


def save_to_sqlite(records: list[dict[str, Any]]) -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)

    with sqlite3.connect(DATABASE_FILE) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS api_records (
                source_id TEXT PRIMARY KEY,
                name TEXT,
                username TEXT,
                email TEXT UNIQUE,
                phone TEXT,
                website TEXT,
                city TEXT,
                company TEXT,
                synced_at TEXT
            )
        """)

        for record in records:
            conn.execute("""
                INSERT OR REPLACE INTO api_records (
                    source_id, name, username, email, phone,
                    website, city, company, synced_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                record["source_id"],
                record["name"],
                record["username"],
                record["email"],
                record["phone"],
                record["website"],
                record["city"],
                record["company"],
                record["synced_at"],
            ))

        conn.commit()


def export_to_csv(records: list[dict[str, Any]]) -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)

    fieldnames = [
        "source_id", "name", "username", "email", "phone",
        "website", "city", "company", "synced_at"
    ]

    with CSV_FILE.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()

        for record in records:
            writer.writerow({field: record.get(field, "") for field in fieldnames})


def export_to_excel(clean_records: list[dict[str, Any]], invalid_records: list[dict[str, Any]]) -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"

    clean_sheet = wb.create_sheet("Clean Records")
    invalid_sheet = wb.create_sheet("Invalid Records")

    create_dashboard_sheet(dashboard, clean_records, invalid_records)
    write_records_sheet(clean_sheet, clean_records, include_error=False)
    write_records_sheet(invalid_sheet, invalid_records, include_error=True)

    wb.save(EXCEL_FILE)


def style_header(ws, row: int = 1) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")

    for cell in ws[row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_records_sheet(ws, records: list[dict[str, Any]], include_error: bool) -> None:
    headers = [
        "source_id", "name", "username", "email", "phone",
        "website", "city", "company", "synced_at"
    ]

    if include_error:
        headers.append("error")

    ws.append(headers)

    for record in records:
        ws.append([record.get(header, "") for header in headers])

    style_header(ws)

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            max_length = max(max_length, len(str(cell.value or "")))

        ws.column_dimensions[column_letter].width = min(max_length + 3, 35)

    ws.freeze_panes = "A2"


def create_dashboard_sheet(ws, clean_records: list[dict[str, Any]], invalid_records: list[dict[str, Any]]) -> None:
    total_raw = len(clean_records) + len(invalid_records)
    total_clean = len(clean_records)
    total_invalid = len(invalid_records)
    unique_companies = len({record["company"] for record in clean_records if record["company"]})
    unique_cities = len({record["city"] for record in clean_records if record["city"]})

    ws["A1"] = "API Data Sync Report"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E78")
    ws.merge_cells("A1:B1")

    rows = [
        ["Metric", "Value"],
        ["Raw records", total_raw],
        ["Clean records", total_clean],
        ["Invalid records", total_invalid],
        ["Unique companies", unique_companies],
        ["Unique cities", unique_cities],
    ]

    for row in rows:
        ws.append(row)

    style_header(ws, row=2)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18

    for row in range(3, 8):
        ws[f"B{row}"].alignment = Alignment(horizontal="center")

    chart = BarChart()
    chart.title = "Sync Summary"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Metric"

    data = Reference(ws, min_col=2, min_row=3, max_row=7)
    categories = Reference(ws, min_col=1, min_row=3, max_row=7)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "D2")


def print_summary(clean_records: list[dict[str, Any]], invalid_records: list[dict[str, Any]]) -> None:
    print("API Data Sync complete")
    print(f"Clean records: {len(clean_records)}")
    print(f"Invalid records: {len(invalid_records)}")
    print(f"SQLite database: {DATABASE_FILE}")
    print(f"CSV export: {CSV_FILE}")
    print(f"Excel report: {EXCEL_FILE}")


def main() -> None:
    # Set use_sample=True if you want the project to work without internet.
    raw_records = fetch_api_data(use_sample=False)
    clean, invalid = clean_records(raw_records)

    save_to_sqlite(clean)
    export_to_csv(clean)
    export_to_excel(clean, invalid)

    print_summary(clean, invalid)


if __name__ == "__main__":
    main()