import csv
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from api_data_sync import clean_records, export_to_csv, export_to_excel, save_to_sqlite


def sample_records():
    return [
        {
            "id": 1,
            "name": " john smith ",
            "username": "john",
            "email": "JOHN@example.com",
            "phone": "123",
            "website": "example.com",
            "address": {"city": "Austin"},
            "company": {"name": "Example LLC"},
        },
        {
            "id": 2,
            "name": "Duplicate Contact",
            "username": "duplicate",
            "email": "john@example.com",
            "phone": "456",
            "website": "example.net",
            "address": {"city": "Dallas"},
            "company": {"name": "Another LLC"},
        },
        {
            "id": 3,
            "name": "Broken Email",
            "username": "broken",
            "email": "not-an-email",
            "phone": "789",
            "website": "example.org",
            "address": {"city": "Houston"},
            "company": {"name": "Broken LLC"},
        },
    ]


def test_clean_records_normalizes_and_validates():
    clean, invalid = clean_records(sample_records())

    assert len(clean) == 1
    assert len(invalid) == 2
    assert clean[0]["name"] == "John Smith"
    assert clean[0]["email"] == "john@example.com"
    assert clean[0]["city"] == "Austin"
    assert any("Duplicate email" in record["error"] for record in invalid)
    assert any("Invalid email" in record["error"] for record in invalid)


def test_sqlite_and_csv_exports(tmp_path: Path):
    clean, _ = clean_records(sample_records())
    database_file = tmp_path / "records.db"
    csv_file = tmp_path / "records.csv"

    save_to_sqlite(clean, database_file)
    export_to_csv(clean, csv_file)

    with sqlite3.connect(database_file) as connection:
        count = connection.execute("SELECT COUNT(*) FROM api_records").fetchone()[0]
    assert count == 1

    with csv_file.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    assert rows[0]["email"] == "john@example.com"


def test_excel_export_contains_expected_sheets(tmp_path: Path):
    clean, invalid = clean_records(sample_records())
    excel_file = tmp_path / "report.xlsx"

    export_to_excel(clean, invalid, excel_file)
    workbook = load_workbook(excel_file)

    assert workbook.sheetnames == ["Dashboard", "Clean Records", "Invalid Records"]
    assert workbook["Clean Records"]["D2"].value == "john@example.com"
